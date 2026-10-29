"""Оформление результата: Markdown → HTML для Telegram и файл Word."""
import html
import io
import re

from docx import Document
from docx.shared import Pt

from i18n import WARN

DOCX = {
    "ru": {"doubt": "Слова, в которых распознавание не уверено: ", "article": "Черновик статьи",
           "transcript": "Расшифровка", "foot": "Конспект сделан бесплатно в Telegram-боте @{bot} — t.me/{bot}?start=docx",
           "foot_text": "Расшифровка сделана бесплатно в Telegram-боте @{bot} — t.me/{bot}?start=docx"},
    "en": {"doubt": "Words speech recognition is unsure about: ", "article": "Draft article",
           "transcript": "Transcript", "foot": "Notes made for free with the Telegram bot @{bot} — t.me/{bot}?start=docx",
           "foot_text": "Transcribed for free with the Telegram bot @{bot} — t.me/{bot}?start=docx"},
}


def _inline(s: str) -> str:
    s = html.escape(s)
    s = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", s)
    s = re.sub(r"(?<![\w*])\*(?!\s)(.+?)(?<!\s)\*(?![\w*])", r"<i>\1</i>", s)
    return s


def to_html(md: str) -> str:
    out = []
    for line in md.splitlines():
        t = line.strip()
        if m := re.match(r"^(#{1,6})\s+(.*)", t):
            level, text = len(m.group(1)), m.group(2).strip("* ")
            out.append(("\n" if out else "") + (f"<b>{html.escape(text.upper())}</b>" if level <= 2
                                                  else f"<b>{html.escape(text)}</b>"))
        elif m := re.match(r"^[-*•]\s+(.*)", t):
            out.append("• " + _inline(m.group(1)))
        elif m := re.match(r"^(\d+)[.)]\s+(.*)", t):
            out.append(f"{m.group(1)}. " + _inline(m.group(2)))
        else:
            out.append(_inline(t))
    return re.sub(r"\n{3,}", "\n\n", "\n".join(out)).strip()


def split_html(text: str, size: int = 3800) -> list[str]:
    """Режет по строкам (теги у нас не переходят через строку)."""
    parts, cur = [], ""
    for line in text.split("\n"):
        while len(line) > size:          # одна огромная строка — режем по пробелу
            cut = line.rfind(" ", 0, size) if " " in line[:size] else size
            parts.append((cur + "\n" + line[:cut]).strip())
            cur, line = "", line[cut:].strip()
        if len(cur) + len(line) + 1 > size:
            parts.append(cur.strip())
            cur = ""
        cur += line + "\n"
    if cur.strip():
        parts.append(cur.strip())
    return parts


def _add_md(doc: Document, md: str, shift: int = 0) -> None:
    for line in md.splitlines():
        t = line.strip()
        if not t:
            continue
        if m := re.match(r"^(#{1,6})\s+(.*)", t):
            doc.add_heading(m.group(2).strip("* "), min(len(m.group(1)) + shift, 9))
            continue
        style = None
        if m := re.match(r"^[-*•]\s+(.*)", t):
            style, t = "List Bullet", m.group(1)
        p = doc.add_paragraph(style=style)
        for i, piece in enumerate(re.split(r"\*\*(.+?)\*\*", t)):
            if piece:
                p.add_run(piece).bold = bool(i % 2)


def make_docx(title: str, summary: str, article: str, transcript: str, doubt: list[str], bot: str = "",
              lang: str = "ru") -> bytes:
    t = DOCX[lang]
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    if bot:   # подпись внизу каждой страницы — так о боте узнают те, кому перешлют файл
        foot = doc.sections[0].footer.paragraphs[0]
        foot.text = t["foot"].format(bot=bot)
        foot.runs[0].font.size = Pt(8)
    doc.add_paragraph(WARN[lang]).runs[0].italic = True
    if doubt:
        doc.add_paragraph(t["doubt"] + ", ".join(doubt)).runs[0].italic = True
    _add_md(doc, summary)
    doc.add_page_break()
    doc.add_heading(t["article"], 1)
    _add_md(doc, article, shift=1)
    doc.add_page_break()
    doc.add_heading(t["transcript"], 1)
    for para in transcript.split("\n\n"):
        doc.add_paragraph(para)
    buf = io.BytesIO()
    doc.core_properties.title = title
    doc.save(buf)
    return buf.getvalue()


def make_transcript_docx(title: str, transcript: str, doubt: list[str], warn: str, bot: str = "",
                         lang: str = "ru") -> bytes:
    """Только дословная расшифровка — без конспекта и статьи."""
    t = DOCX[lang]
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    if bot:
        foot = doc.sections[0].footer.paragraphs[0]
        foot.text = t["foot_text"].format(bot=bot)
        foot.runs[0].font.size = Pt(8)
    doc.add_heading(title, 1)
    doc.add_paragraph(warn).runs[0].italic = True
    if doubt:
        doc.add_paragraph(t["doubt"] + ", ".join(doubt)).runs[0].italic = True
    for para in transcript.split("\n\n"):
        doc.add_paragraph(para)
    buf = io.BytesIO()
    doc.core_properties.title = title
    doc.save(buf)
    return buf.getvalue()


def marked(*texts: str) -> list[str]:
    """Слова, которые нейросеть пометила «(?)» как возможную ошибку распознавания."""
    found = []
    for t in texts:
        for w in re.findall(r"([\w\-]+)[*»\"']{0,3}\s*\((?:\?|возможно,? ошибка распознавания|possibly a recognition error)\)", t, re.I):
            w = w.strip()
            if w.lower() in ("лектором", "лектор", "лектора", "слово", "термин", "текста", "lecturer", "word", "term"):
                continue
            if w.lower() not in (x.lower() for x in found):
                found.append(w)
    return found


def title_of(summary: str) -> str:
    m = re.search(r"^#\s+(.+)$", summary, re.M)
    return (m.group(1).strip("* ") if m else "Конспект")[:80]


def safe_filename(name: str) -> str:
    name = re.sub(r'[\\/:*?"<>|\n\r\t]+', " ", name).strip(" .")
    return (name or "Конспект")[:60]
